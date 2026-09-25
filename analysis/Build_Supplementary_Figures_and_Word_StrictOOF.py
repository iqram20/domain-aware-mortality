#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Build the final supplementary-figure package for the temporal-safe strict-OOF
INSPIRE 30-day postoperative mortality analysis.

Outputs
-------
1. Figure S1:
   Extended model performance/clinical utility 2x2 composite assembled from:
     - ROC_all_models.png
     - PR_all_models.png
     - Calibration_all_models.png
     - Decision_curve_analysis.png

2. Figures S2-S6:
   Canonical copies of the finalized strict-OOF supplementary analyses:
     - FigureS2_Subgroup_Specificity_PPV_StrictOOF.png
     - FigureS3_Tail_Focused_Calibration_StrictOOF.png
     - FigS_Error_Analysis_StrictOOF.png
     - FigS_Leave_One_Domain_Out_Retraining.png
     - FigS_TemporalSafe_StrictOOF_XAI_Input_Correlation_2x2.png

3. Figure S7:
   Eye/Ear sensitivity analysis generated directly from final saved CSV artifacts:
     A. Descriptive mortality prevalence with Wilson 95% CIs
     B. Adjusted odds-ratio forest
     C. Full vs Eye/Ear-feature-removed AUROC/AUPRC
     D. Paired-bootstrap AUROC/AUPRC differences

4. Supplementary_Figures_TemporalSafe_StrictOOF.docx
   One figure per page with manuscript-ready captions.

5. Supplementary_Figure_Manifest.csv

No model is trained, recalibrated, or refit by this script.
All result panels are built from final saved run artifacts.
"""

from __future__ import annotations

import argparse
import math
import pathlib
import shutil
from typing import Dict, List, Mapping, Optional, Sequence, Tuple

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import matplotlib.image as mpimg


DEFAULT_RUN_DIR = pathlib.Path(
    "/athena/madelab/scratch/iqh4001/VitalDB/Inspire/"
    "Results/darn_cv_baseline_runs/"
    "v6_temporal_safe_strict_oof_parallel/"
    "run_20260919_115653"
)

DEFAULT_OUTPUT_SUBDIR = "supplementary_figures_temporal_safe_strict_oof"

FONT_FAMILY = "DejaVu Sans"
BASE_FONT = 13
PANEL_TITLE_FONT = 15
AXIS_LABEL_FONT = 14
TICK_FONT = 12.5
LEGEND_FONT = 11.5
FIGURE_TITLE_FONT = 19

REFERENCE_MODEL = "Hybrid DARN-XGB Blend"
DARN_MODEL = "Uniform DARN"

plt.rcParams.update(
    {
        "font.family": FONT_FAMILY,
        "font.size": BASE_FONT,
        "axes.titlesize": PANEL_TITLE_FONT,
        "axes.labelsize": AXIS_LABEL_FONT,
        "xtick.labelsize": TICK_FONT,
        "ytick.labelsize": TICK_FONT,
        "legend.fontsize": LEGEND_FONT,
        "axes.spines.top": False,
        "axes.spines.right": False,
        "figure.facecolor": "white",
        "savefig.facecolor": "white",
        "pdf.fonttype": 42,
        "ps.fonttype": 42,
    }
)


# =============================================================================
# UTILITIES
# =============================================================================

def require_file(path: pathlib.Path) -> pathlib.Path:
    if not path.exists():
        raise FileNotFoundError(f"Required file not found:\n{path}")
    return path


def find_exact_basename(
    root: pathlib.Path,
    basename: str,
    *,
    excluded_root: Optional[pathlib.Path] = None,
) -> pathlib.Path:
    matches: List[pathlib.Path] = []

    for path in root.rglob(basename):
        if excluded_root is not None:
            try:
                path.relative_to(excluded_root)
                continue
            except ValueError:
                pass
        matches.append(path)

    matches = sorted(set(matches))

    if len(matches) == 0:
        raise FileNotFoundError(
            f"Could not find required figure '{basename}' under:\n{root}"
        )

    if len(matches) > 1:
        raise RuntimeError(
            f"Ambiguous figure basename '{basename}'. Found:\n"
            + "\n".join(f"  - {path}" for path in matches)
        )

    return matches[0]


def save_figure_all_formats(
    fig,
    png_path: pathlib.Path,
    *,
    dpi: int = 300,
) -> None:
    png_path.parent.mkdir(parents=True, exist_ok=True)

    fig.savefig(
        png_path,
        dpi=dpi,
        bbox_inches="tight",
    )
    fig.savefig(
        png_path.with_suffix(".pdf"),
        bbox_inches="tight",
    )
    fig.savefig(
        png_path.with_suffix(".svg"),
        bbox_inches="tight",
    )


def add_panel_label(
    axis,
    label: str,
) -> None:
    axis.text(
        -0.10,
        1.06,
        label,
        transform=axis.transAxes,
        fontsize=16,
        fontweight="bold",
        ha="left",
        va="bottom",
    )


def wilson_interval(
    events: int,
    n: int,
    z: float = 1.959963984540054,
) -> Tuple[float, float]:
    if n <= 0:
        return np.nan, np.nan

    p = events / n
    z2 = z * z
    denom = 1.0 + z2 / n
    center = (p + z2 / (2.0 * n)) / denom
    half = (
        z
        * math.sqrt(
            (p * (1.0 - p) / n)
            + (z2 / (4.0 * n * n))
        )
        / denom
    )

    return max(0.0, center - half), min(1.0, center + half)


def fmt_p(value: float) -> str:
    if not np.isfinite(value):
        return "NA"
    if value < 0.001:
        return "<0.001"
    return f"{value:.3f}"


# =============================================================================
# FIGURE S1 — EXTENDED MODEL PERFORMANCE
# =============================================================================

def build_figure_s1(
    run_dir: pathlib.Path,
    output_dir: pathlib.Path,
    *,
    dpi: int,
) -> pathlib.Path:
    source_names = [
        "ROC_all_models.png",
        "PR_all_models.png",
        "Calibration_all_models.png",
        "Decision_curve_analysis.png",
    ]

    source_paths = [
        find_exact_basename(
            run_dir,
            name,
            excluded_root=output_dir,
        )
        for name in source_names
    ]

    panel_titles = [
        "ROC curves",
        "Precision–recall curves",
        "Calibration",
        "Decision-curve analysis",
    ]

    fig, axes = plt.subplots(
        2,
        2,
        figsize=(18, 15),
    )

    for axis, path, panel_title, panel_label in zip(
        axes.flatten(),
        source_paths,
        panel_titles,
        list("ABCD"),
    ):
        image = mpimg.imread(path)
        axis.imshow(image)
        axis.set_axis_off()
        axis.set_title(
            panel_title,
            fontweight="bold",
            loc="left",
            pad=10,
        )
        axis.text(
            -0.02,
            1.03,
            panel_label,
            transform=axis.transAxes,
            fontsize=17,
            fontweight="bold",
            ha="left",
            va="bottom",
        )

    fig.suptitle(
        "Extended Model Performance and Clinical Utility",
        fontsize=FIGURE_TITLE_FONT,
        fontweight="bold",
        y=0.995,
    )

    fig.tight_layout(
        rect=(0.02, 0.02, 0.99, 0.97),
        h_pad=2.0,
        w_pad=2.0,
    )

    output_path = output_dir / "FigureS1_Extended_Model_Performance_2x2.png"
    save_figure_all_formats(fig, output_path, dpi=dpi)
    plt.close(fig)

    return output_path


# =============================================================================
# FIGURES S2-S6 — FINALIZED EXISTING STRICT-OOF FIGURES
# =============================================================================

CANONICAL_EXISTING_SUPPLEMENTARY = {
    "Figure S2": "FigureS2_Subgroup_Specificity_PPV_StrictOOF.png",
    "Figure S3": "FigureS3_Tail_Focused_Calibration_StrictOOF.png",
    "Figure S4": "FigS_Error_Analysis_StrictOOF.png",
    "Figure S5": "FigS_Leave_One_Domain_Out_Retraining.png",
    "Figure S6": "FigS_TemporalSafe_StrictOOF_XAI_Input_Correlation_2x2.png",
}


def collect_existing_supplementary(
    run_dir: pathlib.Path,
    output_dir: pathlib.Path,
) -> Dict[str, pathlib.Path]:
    resolved: Dict[str, pathlib.Path] = {}

    for figure_id, basename in CANONICAL_EXISTING_SUPPLEMENTARY.items():
        source = find_exact_basename(
            run_dir,
            basename,
            excluded_root=output_dir,
        )

        destination = output_dir / basename
        shutil.copy2(source, destination)

        # Copy matching vector/PDF formats if available next to source.
        for suffix in [".pdf", ".svg"]:
            companion = source.with_suffix(suffix)
            if companion.exists():
                shutil.copy2(
                    companion,
                    output_dir / companion.name,
                )

        resolved[figure_id] = destination

    return resolved


# =============================================================================
# FIGURE S7 — EYE/EAR SENSITIVITY
# =============================================================================

def load_eye_ear_sources(
    run_dir: pathlib.Path,
) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    adjusted = pd.read_csv(
        require_file(
            run_dir / "eye_ear_adjusted_association.csv"
        )
    )

    feature_dir = run_dir / "eye_ear_feature_retraining"

    removed_metrics = pd.read_csv(
        require_file(
            feature_dir / "eye_ear_feature_removal_test_metrics.csv"
        )
    )

    paired = pd.read_csv(
        require_file(
            feature_dir / "eye_ear_feature_removal_paired_bootstrap.csv"
        )
    )

    full_metrics = pd.read_csv(
        require_file(
            run_dir / "test_model_metrics.csv"
        )
    )

    return adjusted, removed_metrics, paired, full_metrics


def build_figure_s7(
    run_dir: pathlib.Path,
    output_dir: pathlib.Path,
    *,
    dpi: int,
) -> pathlib.Path:
    adjusted, removed_metrics, paired, full_metrics = load_eye_ear_sources(
        run_dir
    )

    fig, axes = plt.subplots(
        2,
        2,
        figsize=(17, 14),
    )
    ax_a, ax_b, ax_c, ax_d = axes.flatten()

    # -------------------------------------------------------------------------
    # A. Descriptive mortality prevalence
    # -------------------------------------------------------------------------
    descriptive = adjusted.loc[
        adjusted["analysis"].astype(str).eq("Descriptive mortality")
    ].copy()

    if descriptive.empty:
        raise ValueError(
            "eye_ear_adjusted_association.csv has no descriptive mortality rows."
        )

    labels: List[str] = []
    rates: List[float] = []
    yerr_low: List[float] = []
    yerr_high: List[float] = []

    for _, row in descriptive.iterrows():
        label = str(row["model_stage"]).replace("Eye/ear ", "").title()
        n = int(row["N"])
        deaths = int(row["Deaths"])
        rate = deaths / n
        low, high = wilson_interval(deaths, n)

        labels.append(label)
        rates.append(100.0 * rate)
        yerr_low.append(100.0 * (rate - low))
        yerr_high.append(100.0 * (high - rate))

    x = np.arange(len(labels))
    ax_a.bar(
        x,
        rates,
        width=0.62,
    )
    ax_a.errorbar(
        x,
        rates,
        yerr=np.vstack([yerr_low, yerr_high]),
        fmt="none",
        capsize=5,
        linewidth=1.3,
    )
    ax_a.set_xticks(x)
    ax_a.set_xticklabels(labels)
    ax_a.set_ylabel("30-day mortality (%)")
    ax_a.set_title(
        "Mortality by Eye/Ear diagnosis status",
        fontweight="bold",
        loc="left",
    )
    ax_a.grid(axis="y", linestyle="--", alpha=0.20)

    for xx, yy, (_, row) in zip(x, rates, descriptive.iterrows()):
        ax_a.text(
            xx,
            yy,
            f"{yy:.2f}%\n({int(row['Deaths'])}/{int(row['N']):,})",
            ha="center",
            va="bottom",
            fontsize=10.5,
        )

    # -------------------------------------------------------------------------
    # B. Adjusted odds-ratio forest
    # -------------------------------------------------------------------------
    or_rows = adjusted.loc[
        pd.to_numeric(
            adjusted["odds_ratio"],
            errors="coerce",
        ).notna()
    ].copy()

    if or_rows.empty:
        raise ValueError("No adjusted OR rows found.")

    stage_labels = (
        or_rows["model_stage"]
        .astype(str)
        .str.replace(
            "Additionally adjusted for surgical department",
            "Age + sex + ASA + urgency + department",
            regex=False,
        )
        .str.replace(
            "Age, sex, ASA, and urgency adjusted",
            "Age + sex + ASA + urgency",
            regex=False,
        )
        .str.replace(
            "Age and sex adjusted",
            "Age + sex",
            regex=False,
        )
        .tolist()
    )

    odds = pd.to_numeric(or_rows["odds_ratio"], errors="coerce").to_numpy()
    low = pd.to_numeric(or_rows["CI_low"], errors="coerce").to_numpy()
    high = pd.to_numeric(or_rows["CI_high"], errors="coerce").to_numpy()
    p_values = pd.to_numeric(or_rows["p_value"], errors="coerce").to_numpy()

    y = np.arange(len(or_rows))[::-1]

    ax_b.errorbar(
        odds,
        y,
        xerr=np.vstack([odds - low, high - odds]),
        fmt="o",
        capsize=4,
        linewidth=1.4,
        markersize=7,
    )
    ax_b.axvline(
        1.0,
        linestyle="--",
        linewidth=1.2,
    )
    ax_b.set_xscale("log")
    ax_b.set_yticks(y)
    ax_b.set_yticklabels(stage_labels)
    ax_b.set_xlabel("Odds ratio for Eye/Ear diagnosis (95% CI)")
    ax_b.set_title(
        "Adjusted association with 30-day mortality",
        fontweight="bold",
        loc="left",
    )
    ax_b.grid(axis="x", linestyle="--", alpha=0.20)

    for yy, estimate, upper, p_value in zip(
        y,
        odds,
        high,
        p_values,
    ):
        ax_b.text(
            upper * 1.04,
            yy,
            f"p={fmt_p(p_value)}",
            va="center",
            fontsize=10,
        )

    # -------------------------------------------------------------------------
    # C. Full vs Eye/Ear-feature-removed discrimination
    # -------------------------------------------------------------------------
    models = [DARN_MODEL, REFERENCE_MODEL]
    metrics = ["AUROC", "AUPRC"]

    full = full_metrics.loc[
        full_metrics["model"].astype(str).isin(models)
    ].copy()
    removed = removed_metrics.loc[
        removed_metrics["model"].astype(str).isin(models)
    ].copy()

    rows_c: List[Dict[str, object]] = []

    for model in models:
        full_row = full.loc[
            full["model"].astype(str).eq(model)
        ]
        removed_row = removed.loc[
            removed["model"].astype(str).eq(model)
        ]

        if full_row.empty or removed_row.empty:
            raise ValueError(
                f"Missing full/removed metric row for model: {model}"
            )

        for metric in metrics:
            rows_c.append(
                {
                    "model": model,
                    "metric": metric,
                    "full": float(full_row.iloc[0][metric]),
                    "removed": float(removed_row.iloc[0][metric]),
                }
            )

    comp = pd.DataFrame(rows_c)

    group_centers = np.arange(len(models))
    width = 0.18

    for metric_index, metric in enumerate(metrics):
        subset = comp.loc[comp["metric"].eq(metric)].copy()

        offset = (
            -0.27
            if metric == "AUROC"
            else 0.27
        )

        ax_c.bar(
            group_centers + offset - width / 2,
            subset["full"],
            width=width,
            label=f"{metric}: full",
        )
        ax_c.bar(
            group_centers + offset + width / 2,
            subset["removed"],
            width=width,
            label=f"{metric}: Eye/Ear removed",
        )

    ax_c.set_xticks(group_centers)
    ax_c.set_xticklabels(
        ["Uniform DARN", "Hybrid Blend"],
    )
    ax_c.set_ylim(0.0, 1.02)
    ax_c.set_ylabel("Held-out test metric")
    ax_c.set_title(
        "Discrimination before and after Eye/Ear feature removal",
        fontweight="bold",
        loc="left",
    )
    ax_c.legend(
        frameon=False,
        fontsize=9.5,
        ncol=2,
    )
    ax_c.grid(axis="y", linestyle="--", alpha=0.20)

    # -------------------------------------------------------------------------
    # D. Paired-bootstrap differences
    # -------------------------------------------------------------------------
    paired_plot = paired.loc[
        paired["model"].astype(str).isin(models)
        & paired["metric"].astype(str).isin(metrics)
    ].copy()

    if paired_plot.empty:
        raise ValueError("No Eye/Ear paired AUROC/AUPRC rows found.")

    paired_plot["label"] = (
        paired_plot["model"]
        .replace(
            {
                DARN_MODEL: "Uniform DARN",
                REFERENCE_MODEL: "Hybrid Blend",
            }
        )
        .astype(str)
        + " — "
        + paired_plot["metric"].astype(str)
    )

    paired_plot = paired_plot.reset_index(drop=True)
    y_d = np.arange(len(paired_plot))[::-1]

    diff = pd.to_numeric(
        paired_plot["difference_favoring_complete_model"],
        errors="coerce",
    ).to_numpy()
    ci_low = pd.to_numeric(
        paired_plot["CI_low"],
        errors="coerce",
    ).to_numpy()
    ci_high = pd.to_numeric(
        paired_plot["CI_high"],
        errors="coerce",
    ).to_numpy()
    p_d = pd.to_numeric(
        paired_plot["bootstrap_p_two_sided"],
        errors="coerce",
    ).to_numpy()

    ax_d.errorbar(
        diff,
        y_d,
        xerr=np.vstack(
            [
                diff - ci_low,
                ci_high - diff,
            ]
        ),
        fmt="o",
        capsize=4,
        linewidth=1.4,
        markersize=7,
    )
    ax_d.axvline(
        0.0,
        linestyle="--",
        linewidth=1.2,
    )
    ax_d.set_yticks(y_d)
    ax_d.set_yticklabels(paired_plot["label"])
    ax_d.set_xlabel(
        "Full model − Eye/Ear-removed model\n(positive favors full model)"
    )
    ax_d.set_title(
        "Paired-bootstrap discrimination differences",
        fontweight="bold",
        loc="left",
    )
    ax_d.grid(axis="x", linestyle="--", alpha=0.20)

    for yy, high_ci, p_value in zip(
        y_d,
        ci_high,
        p_d,
    ):
        ax_d.text(
            high_ci + 0.002,
            yy,
            f"p={fmt_p(p_value)}",
            va="center",
            fontsize=10,
        )

    # Panel labels and title
    for axis, label in zip(
        axes.flatten(),
        list("ABCD"),
    ):
        add_panel_label(axis, label)

    fig.suptitle(
        "Eye/Ear Diagnosis Sensitivity Analysis",
        fontsize=FIGURE_TITLE_FONT,
        fontweight="bold",
        y=0.995,
    )

    fig.tight_layout(
        rect=(0.03, 0.03, 0.99, 0.97),
        h_pad=3.2,
        w_pad=3.4,
    )

    output_path = output_dir / "FigureS7_EyeEar_Sensitivity_2x2.png"
    save_figure_all_formats(fig, output_path, dpi=dpi)
    plt.close(fig)

    return output_path


# =============================================================================
# CAPTIONS
# =============================================================================

def build_captions(
    run_dir: pathlib.Path,
) -> Dict[str, str]:
    metrics = pd.read_csv(
        require_file(
            run_dir / "test_model_metrics.csv"
        )
    )

    n_value = None
    event_value = None

    for candidate in ["N", "n"]:
        if candidate in metrics.columns:
            values = pd.to_numeric(metrics[candidate], errors="coerce").dropna()
            if len(values):
                n_value = int(values.iloc[0])
                break

    for candidate in ["Events", "events"]:
        if candidate in metrics.columns:
            values = pd.to_numeric(metrics[candidate], errors="coerce").dropna()
            if len(values):
                event_value = int(values.iloc[0])
                break

    cohort_phrase = ""
    if n_value is not None and event_value is not None:
        cohort_phrase = (
            f" Analyses were evaluated in the held-out test cohort "
            f"(N={n_value:,}; deaths={event_value:,})."
        )

    return {
        "Figure S1": (
            "Figure S1. Extended model performance and clinical utility. "
            "(A) Receiver operating characteristic curves for all saved models. "
            "(B) Precision–recall curves for all saved models. "
            "(C) Calibration curves comparing predicted and observed 30-day "
            "postoperative mortality risk. "
            "(D) Decision-curve analysis across clinically relevant threshold "
            "probabilities."
            + cohort_phrase
        ),
        "Figure S2": (
            "Figure S2. Exploratory subgroup operating-point performance. "
            "Specificity and positive predictive value are shown across the "
            "prespecified clinical subgroups at the locked development-defined "
            "operating threshold. Subgroup estimates are descriptive and should "
            "be interpreted cautiously where event counts are small."
        ),
        "Figure S3": (
            "Figure S3. Tail-focused calibration of postoperative mortality risk. "
            "Observed outcomes are compared with model-predicted risk within the "
            "high-risk tail and risk-stratified groups, complementing the overall "
            "calibration analyses shown in the main performance figure."
        ),
        "Figure S4": (
            "Figure S4. Clinical error analysis at the locked operating threshold. "
            "Descriptive characteristics of true-positive, false-positive, "
            "false-negative, and true-negative predictions are summarized to "
            "support model auditing. These comparisons are descriptive rather "
            "than causal."
        ),
        "Figure S5": (
            "Figure S5. Leave-one-domain-out retraining analysis. "
            "Model performance after complete retraining with one clinical domain "
            "removed is compared with the corresponding complete-model performance. "
            "This analysis evaluates domain-level robustness and redundancy and "
            "complements the fitted-model input-domain masking analysis."
        ),
        "Figure S6": (
            "Figure S6. Input-correlation and explainability sensitivity analysis. "
            "Relationships among important perioperative inputs and component-level "
            "feature attributions are shown to characterize correlation structure "
            "and the stability of feature-level explanations across the DARN and "
            "XGBoost model components."
        ),
        "Figure S7": (
            "Figure S7. Eye/Ear diagnosis sensitivity analysis. "
            "(A) Descriptive 30-day mortality prevalence according to the presence "
            "or absence of the Eye/Ear diagnosis indicator, with Wilson 95% "
            "confidence intervals. "
            "(B) Odds ratios for the Eye/Ear diagnosis indicator across sequential "
            "adjustment models. "
            "(C) Held-out AUROC and AUPRC for the complete models and models "
            "retrained after removing the Eye/Ear diagnosis feature. "
            "(D) Paired-bootstrap differences in AUROC and AUPRC; positive values "
            "favor the complete model. The feature-removal analysis is a "
            "sensitivity analysis and is not interpreted causally."
        ),
    }


# =============================================================================
# WORD COMPILATION
# =============================================================================

def build_word_document(
    figure_paths: Mapping[str, pathlib.Path],
    captions: Mapping[str, str],
    output_path: pathlib.Path,
    run_dir: pathlib.Path,
) -> pathlib.Path:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt
    except ImportError as exc:
        raise ImportError(
            "python-docx is required. Install it in the active environment."
        ) from exc

    order = [
        "Figure S1",
        "Figure S2",
        "Figure S3",
        "Figure S4",
        "Figure S5",
        "Figure S6",
        "Figure S7",
    ]

    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    title = document.add_heading(
        "Supplementary Figures — Temporal-Safe Strict-OOF Analysis",
        level=1,
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    provenance = document.add_paragraph(
        f"Final analysis run: {run_dir}"
    )
    provenance.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_paragraph(
        "All figures are derived from saved final-run artifacts. "
        "No model training, recalibration, or refitting is performed by the "
        "supplementary-figure compilation script."
    )

    for index, figure_id in enumerate(order):
        if index > 0:
            document.add_page_break()

        heading = document.add_heading(
            figure_id,
            level=2,
        )
        heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

        path = require_file(figure_paths[figure_id])

        image_paragraph = document.add_paragraph()
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_run = image_paragraph.add_run()
        image_run.add_picture(
            str(path),
            width=Inches(6.9),
        )

        caption_paragraph = document.add_paragraph()
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        caption_run = caption_paragraph.add_run(
            captions[figure_id]
        )
        caption_run.font.name = "Arial"
        caption_run.font.size = Pt(10)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)

    return output_path


# =============================================================================
# MANIFEST
# =============================================================================

def build_manifest(
    figure_paths: Mapping[str, pathlib.Path],
    captions: Mapping[str, str],
    output_dir: pathlib.Path,
) -> pathlib.Path:
    rows = []

    for order_index, figure_id in enumerate(
        [
            "Figure S1",
            "Figure S2",
            "Figure S3",
            "Figure S4",
            "Figure S5",
            "Figure S6",
            "Figure S7",
        ],
        start=1,
    ):
        path = figure_paths[figure_id]

        rows.append(
            {
                "order": order_index,
                "figure_id": figure_id,
                "filename": path.name,
                "file": str(path),
                "caption": captions[figure_id],
                "status": "created",
            }
        )

    manifest = pd.DataFrame(rows)
    manifest_path = output_dir / "Supplementary_Figure_Manifest.csv"
    manifest.to_csv(manifest_path, index=False)

    return manifest_path


# =============================================================================
# MAIN
# =============================================================================

def run(
    run_dir: pathlib.Path,
    output_dir: pathlib.Path,
    *,
    dpi: int,
) -> None:
    run_dir = run_dir.resolve()
    output_dir = output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    print("=" * 100)
    print("BUILDING SUPPLEMENTARY FIGURES — TEMPORAL-SAFE STRICT-OOF")
    print("=" * 100)
    print("Run directory :", run_dir)
    print("Output        :", output_dir)

    figure_paths: Dict[str, pathlib.Path] = {}

    print("\nBuilding Figure S1...")
    figure_paths["Figure S1"] = build_figure_s1(
        run_dir,
        output_dir,
        dpi=dpi,
    )

    print("Collecting Figures S2-S6...")
    figure_paths.update(
        collect_existing_supplementary(
            run_dir,
            output_dir,
        )
    )

    print("Building Figure S7...")
    figure_paths["Figure S7"] = build_figure_s7(
        run_dir,
        output_dir,        dpi=dpi,
    )

    captions = build_captions(run_dir)

    word_path = output_dir / "Supplementary_Figures_TemporalSafe_StrictOOF.docx"
    build_word_document(
        figure_paths,
        captions,
        word_path,
        run_dir,
    )

    manifest_path = build_manifest(
        figure_paths,
        captions,
        output_dir,
    )

    print("\n" + "=" * 100)
    print("COMPLETE")
    print("=" * 100)
    print("Supplementary figures: 7")
    print("Word file            :", word_path)
    print("Manifest             :", manifest_path)

    for figure_id in [
        "Figure S1",
        "Figure S2",
        "Figure S3",
        "Figure S4",
        "Figure S5",
        "Figure S6",
        "Figure S7",
    ]:
        print(f"{figure_id:10s}: {figure_paths[figure_id]}")


def main() -> None:
    parser = argparse.ArgumentParser()

    parser.add_argument(
        "--run_dir",
        type=pathlib.Path,
        default=DEFAULT_RUN_DIR,
    )
    parser.add_argument(
        "--output_dir",
        type=pathlib.Path,
        default=None,
    )
    parser.add_argument(
        "--dpi",
        type=int,
        default=300,
    )

    args = parser.parse_args()

    output_dir = (
        args.output_dir
        if args.output_dir is not None
        else args.run_dir / DEFAULT_OUTPUT_SUBDIR
    )

    run(
        args.run_dir,
        output_dir,
        dpi=args.dpi,
    )


if __name__ == "__main__":
    main()